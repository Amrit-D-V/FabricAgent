from __future__ import annotations

import io
import re

_MAX_MEASURE_ROWS = 60
_MAX_REL_ROWS = 40


def _parse(tmdl_files: dict[str, str]) -> dict:
    tables = []
    for path, text in tmdl_files.items():
        if not path.startswith("tables/"):
            continue
        tm = re.search(r"^table\s+(?:'([^']+)'|(\S+))", text, re.MULTILINE)
        tname = (tm.group(1) or tm.group(2)) if tm else path
        measures = []
        for mm in re.finditer(
            r"^\s*measure\s+(?:'([^']+)'|([^=\s]+))\s*=(.*?)(?=\n\s*(?:measure|column|partition|hierarchy)\b|\Z)",
            text,
            re.MULTILINE | re.DOTALL,
        ):
            name = (mm.group(1) or mm.group(2) or "").strip()
            fmt = re.search(r"formatString:\s*(.+)", mm.group(0))
            measures.append((name, fmt.group(1).strip() if fmt else ""))
        if measures or tm:
            tables.append({"name": tname, "measures": measures})

    rels = []
    rel_text = tmdl_files.get("relationships.tmdl", "")
    for block in re.finditer(
        r"relationship\s+\S+.*?(?=\nrelationship\s|\Z)", rel_text, re.DOTALL
    ):
        b = block.group(0)
        frm = re.search(r"fromColumn:\s*([^\n]+)", b)
        to = re.search(r"toColumn:\s*([^\n]+)", b)
        if frm and to:
            rels.append((frm.group(1).strip(), to.group(1).strip()))
    return {"tables": tables, "relationships": rels}


def _summary_line(parsed: dict) -> str:
    n_t = len(parsed["tables"])
    n_m = sum(len(t["measures"]) for t in parsed["tables"])
    n_r = len(parsed["relationships"])
    return f"{n_t} tables · {n_m} measures · {n_r} relationships"


def build_pptx(
    tmdl_files: dict[str, str], model_name: str, coverage: dict | None = None
) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    parsed = _parse(tmdl_files)
    prs = Presentation()
    title_only = prs.slide_layouts[5]

    def _title_slide():
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = model_name
        s.placeholders[1].text = _summary_line(parsed)

    def _bullets(title: str, lines: list[str]):
        s = prs.slides.add_slide(title_only)
        s.shapes.title.text = title
        box = s.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(9), Inches(5.5))
        tf = box.text_frame
        tf.word_wrap = True
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = ln
            p.font.size = Pt(14)

    _title_slide()

    if coverage:
        s = coverage["summary"]
        miss = [
            it["requested"] for it in coverage["items"] if it["status"] == "missing"
        ]
        lines = [
            (f"Coverage: {s['coverage_pct']}%   "
            f"(covered {s['covered']} · partial {s['partial']} · missing {s['missing']})"),
            "",
            (
                "Requested items with NO backing measure:"
                if miss
                else "All requested items are backed by a measure."
            ),
        ] + [f"  • {m}" for m in miss[:18]]
        _bullets("Wireframe coverage", lines)

    mlines, shown = [], 0
    for t in parsed["tables"]:
        if not t["measures"]:
            continue
        mlines.append(f"{t['name']}")
        for name, fmt in t["measures"]:
            if shown >= _MAX_MEASURE_ROWS:
                break
            mlines.append(f"   • {name}" + (f"   [{fmt}]" if fmt else ""))
            shown += 1
        if shown >= _MAX_MEASURE_ROWS:
            mlines.append(
                f"   … (+{sum(len(x['measures']) for x in parsed['tables']) - shown} more measures)"
            )
            break
    _bullets("Measures", mlines or ["(no measures)"])

    rels = parsed["relationships"]
    rlines = [f"{frm}  →  {to}" for frm, to in rels[:_MAX_REL_ROWS]]
    if len(rels) > _MAX_REL_ROWS:
        rlines.append(f"… (+{len(rels) - _MAX_REL_ROWS} more)")
    _bullets("Relationships", rlines or ["(no relationships)"])

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_pdf(
    tmdl_files: dict[str, str], model_name: str, coverage: dict | None = None
) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    parsed = _parse(tmdl_files)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, title=f"{model_name} — Model Report")
    styles = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=9, leading=12)
    story = [
        Paragraph(f"Model Report — {model_name}", styles["Title"]),
        Paragraph(_summary_line(parsed), styles["Normal"]),
        Spacer(1, 12),
    ]

    if coverage:
        s = coverage["summary"]
        story.append(Paragraph("Wireframe coverage", styles["Heading2"]))
        story.append(
            Paragraph(
                f"{s['coverage_pct']}% — covered {s['covered']}, partial {s['partial']}, missing {s['missing']}",
                styles["Normal"],
            )
        )
        for it in coverage["items"]:
            if it["status"] == "missing":
                story.append(Paragraph(f"✗ missing: {it['requested']}", small))
        story.append(Spacer(1, 12))

    story.append(Paragraph("Measures", styles["Heading2"]))
    shown = 0
    for t in parsed["tables"]:
        if not t["measures"]:
            continue
        story.append(Paragraph(f"<b>{t['name']}</b>", small))
        for name, fmt in t["measures"]:
            if shown >= _MAX_MEASURE_ROWS:
                break
            story.append(
                Paragraph(f"&bull; {name}" + (f" [{fmt}]" if fmt else ""), small)
            )
            shown += 1
        if shown >= _MAX_MEASURE_ROWS:
            story.append(Paragraph("… (truncated)", small))
            break

    story.append(Spacer(1, 12))
    story.append(Paragraph("Relationships", styles["Heading2"]))
    for frm, to in parsed["relationships"][:_MAX_REL_ROWS]:
        story.append(Paragraph(f"{frm} &rarr; {to}", small))

    doc.build(story)
    return buf.getvalue()
