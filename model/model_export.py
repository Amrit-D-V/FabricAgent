from __future__ import annotations

import glob
import os

from model.model_docs import _parse_relationships, _parse_table, _read
from model.model_quality import find_model_dir


def export_excel(model_dir: str, out_path: str | None = None) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    resolved = find_model_dir(model_dir) or model_dir
    defn = os.path.join(resolved, "definition")
    base = defn if os.path.isdir(defn) else resolved

    table_files = sorted(glob.glob(os.path.join(base, "tables", "*.tmdl"))) or sorted(
        glob.glob(os.path.join(base, "**", "*.tmdl"), recursive=True)
    )
    tables = [_parse_table(_read(f)) for f in table_files]
    tables = [t for t in tables if t["name"]]
    rels = _parse_relationships(_read(os.path.join(base, "relationships.tmdl")))

    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Columns"
    ws.append(["Table", "Column", "Type"])
    for cell in ws[1]:
        cell.font = bold
    for t in tables:
        for c, dt in t["columns"]:
            ws.append([t["name"], c, dt or ""])

    wm = wb.create_sheet("Measures")
    wm.append(["Table", "Measure", "Format", "Expression", "Description"])
    for cell in wm[1]:
        cell.font = bold
    for t in tables:
        for n, e, fmt, desc in t["measures"]:
            wm.append([t["name"], n, fmt, e, desc])

    wr = wb.create_sheet("Relationships")
    wr.append(["From", "To", "Cardinality"])
    for cell in wr[1]:
        cell.font = bold
    for frm, to, card in rels:
        wr.append([frm, to, card])

    out = out_path or os.path.join(resolved, "model_export.xlsx")
    wb.save(out)
    return out
